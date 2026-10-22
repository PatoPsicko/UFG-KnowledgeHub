"""Extracción de texto de los archivos subidos.

La usan el índice del asistente (para buscar dentro de los documentos) y el
vigilante de seguridad (para revisar el contenido antes de publicarlo).
"""

from __future__ import annotations

import os
import re
import zipfile
from xml.etree import ElementTree

MAX_CHARS_PER_FILE = 400_000
MAX_PDF_PAGES = 300
# Un PDF con menos caracteres por página que esto se considera escaneado.
MIN_CHARS_PDF_PAGE = 25


def extraer_paginas(path: str | None) -> tuple[list[tuple[int, str]], str]:
    """Devuelve (lista de (página, texto), estado).

    estado: 'ok', 'sin_texto' (PDF escaneado), 'no_soportado', 'sin_archivo'
    o 'error: ...'. Los formatos sin páginas usan la página 0.
    """
    if not path or not os.path.exists(path):
        return [], 'sin_archivo'

    ext = path.rsplit('.', 1)[-1].lower() if '.' in path else ''
    try:
        if ext in ('txt', 'md', 'csv'):
            with open(path, 'r', encoding='utf-8', errors='replace') as fh:
                return [(0, fh.read(MAX_CHARS_PER_FILE))], 'ok'

        if ext == 'pdf':
            from pypdf import PdfReader
            paginas = []
            total = 0
            lector = PdfReader(path)
            for n, pagina in enumerate(lector.pages[:MAX_PDF_PAGES], 1):
                texto = pagina.extract_text() or ''
                paginas.append((n, texto))
                total += len(texto.strip())
                if total > MAX_CHARS_PER_FILE:
                    break
            # PDF escaneado sin capa de texto: se marca; el OCR queda fuera de alcance.
            if paginas and total < MIN_CHARS_PDF_PAGE * len(paginas):
                return [], 'sin_texto'
            return paginas, 'ok'

        if ext == 'docx':
            import docx
            documento = docx.Document(path)
            partes = [p.text for p in documento.paragraphs]
            for tabla in documento.tables:
                for fila in tabla.rows:
                    partes.append(' | '.join(c.text for c in fila.cells))
            return [(0, '\n'.join(partes)[:MAX_CHARS_PER_FILE])], 'ok'

        if ext in ('pptx', 'xlsx', 'odt', 'odp', 'ods'):
            return _extraer_ooxml(path, ext), 'ok'

        if ext == 'rtf':
            with open(path, 'r', encoding='latin-1', errors='replace') as fh:
                crudo = fh.read(MAX_CHARS_PER_FILE)
            # Quita grupos y palabras de control: aproximación suficiente para buscar.
            texto = re.sub(r'\\[a-z]+-?\d* ?|[{}]', ' ', crudo)
            return [(0, re.sub(r'\s+', ' ', texto))], 'ok'

    except ImportError as exc:
        return [], f'error: falta la librería {exc.name}'
    except Exception as exc:
        return [], f'error: {exc.__class__.__name__}'

    return [], 'no_soportado'


def texto_plano(path: str | None, limite: int = 12000) -> tuple[str, str]:
    """Primeros `limite` caracteres del documento, con el estado de extracción."""
    paginas, estado = extraer_paginas(path)
    texto = '\n'.join(t for _, t in paginas)
    return re.sub(r'[ \t]+', ' ', texto).strip()[:limite], estado


def _textos_xml(datos: bytes, etiquetas: set[str]) -> str:
    """Texto de los nodos cuya etiqueta local está en `etiquetas`, en orden."""
    partes = []
    for nodo in ElementTree.fromstring(datos).iter():
        if nodo.tag.rsplit('}', 1)[-1] in etiquetas:
            texto = ''.join(nodo.itertext()).strip()
            if texto:
                partes.append(texto)
    return '\n'.join(partes)


def _extraer_ooxml(path: str, ext: str) -> list[tuple[int, str]]:
    """PowerPoint, Excel y OpenDocument son ZIP con XML dentro: se leen con la
    biblioteca estándar. En PowerPoint cada diapositiva cuenta como página."""
    with zipfile.ZipFile(path) as z:
        nombres = z.namelist()
        if ext == 'pptx':
            def numero(nombre):
                return int(re.search(r'(\d+)\.xml$', nombre).group(1))
            diapositivas = sorted((n for n in nombres if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                                  key=numero)
            return [(numero(n), _textos_xml(z.read(n), {'p'})[:MAX_CHARS_PER_FILE])
                    for n in diapositivas[:MAX_PDF_PAGES]]
        if ext == 'xlsx':
            partes = []
            if 'xl/sharedStrings.xml' in nombres:
                partes.append(_textos_xml(z.read('xl/sharedStrings.xml'), {'si'}))
            for n in sorted(n for n in nombres if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', n)):
                for celda in ElementTree.fromstring(z.read(n)).iter():
                    # t="s": el valor es un índice a sharedStrings, ya leído arriba.
                    if celda.tag.rsplit('}', 1)[-1] == 'c' and celda.get('t') != 's':
                        texto = ''.join(celda.itertext()).strip()
                        if texto:
                            partes.append(texto)
            return [(0, '\n'.join(partes)[:MAX_CHARS_PER_FILE])]
        # OpenDocument (odt, odp, ods): párrafos y encabezados de content.xml.
        return [(0, _textos_xml(z.read('content.xml'), {'p', 'h'})[:MAX_CHARS_PER_FILE])]
