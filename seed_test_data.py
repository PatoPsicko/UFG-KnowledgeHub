"""Datos de prueba para UFG Knowledge Hub.

Crea, a través de la API (con la app en marcha), lo necesario para probar la
plataforma y el asistente de IA:

    - 8 usuarios de prueba con carrera y semestre
    - 100 documentos con muy poco texto (pdf, docx, pptx, xlsx, txt, md, csv):
        66 en la Biblioteca, 20 en grupos de trabajo y 14 en anotaciones privadas
    - 5 grupos de trabajo con miembros, tareas y mensajes
    - anotaciones de texto, contactos, mensajes directos y chat global
    - solicitudes de contacto pendientes para el usuario demo (demo@ufg.edu.sv)

Cada documento contiene algún dato concreto (fechas, aulas, cifras) para poder
comprobar con el chatbot que lo encuentra y que respeta los permisos: los
grupos "Emprendimiento UFG" y "Club de Psicología Cognitiva" y las anotaciones
privadas NO son visibles para el usuario demo.

Uso:
    python app.py                       # en otra terminal
    python seed_test_data.py            # crea los datos
    python seed_test_data.py --url http://127.0.0.1:5001

Para generar los .pptx y .xlsx necesita python-pptx y openpyxl (sólo este
script; la app no los usa):
    pip install python-pptx openpyxl

Si los usuarios de prueba ya existen, no hace nada (para no duplicar datos).
Para empezar de cero: python manage.py db --reset

Las cuentas de prueba usan el dominio reservado @prueba.ufg.edu.sv, que el
registro normal no acepta (Fase 2: sólo @ufg.edu.sv). Por eso este script las
crea directamente en la base de datos de esta PC y después entra por la API
como cualquier estudiante: tiene que correr en la misma PC que el servidor.
"""

import argparse
import csv
import io
import re
import shutil
import sys
import tempfile
import textwrap
import unicodedata
from pathlib import Path

import requests

# Contraseña común de todas las cuentas de prueba. Son cuentas locales de
# desarrollo; no reutilizar en ningún entorno real.
TEST_PASSWORD = 'PruebaUFG2026'
DEMO_EMAIL = 'demo@ufg.edu.sv'

USUARIOS = [
    # clave, nombre, carrera, semestre, color
    ('ana', 'Ana Martínez', 'Ingeniería en Sistemas', 6, '#3B82F6'),
    ('carlos', 'Carlos Hernández', 'Ingeniería en Sistemas', 8, '#10B981'),
    ('lucia', 'Lucía Ramírez', 'Arquitectura', 4, '#F59E0B'),
    ('diego', 'Diego Flores', 'Administración de Empresas', 5, '#EF4444'),
    ('sofia', 'Sofía Castillo', 'Psicología', 3, '#8B5CF6'),
    ('javier', 'Javier Morales', 'Ingeniería Civil', 7, '#EC4899'),
    ('valeria', 'Valeria Romero', 'Medicina', 2, '#14B8A6'),
    ('andres', 'Andrés López', 'Derecho', 9, '#6366F1'),
]


def email_de(clave):
    nombre = next(u[1] for u in USUARIOS if u[0] == clave)
    return slug(nombre, '.') + '@prueba.ufg.edu.sv'


# --- Contenido -----------------------------------------------------------------
# Biblioteca: (categoría, autor, etiquetas, [(título, texto), x3])
BIBLIOTECA = [
    ('Base de Datos', 'Ing. Ricardo Menjívar', 'sql, normalización', [
        ('Apuntes de normalización',
         'La primera forma normal (1FN) exige valores atómicos. La segunda (2FN) elimina dependencias '
         'parciales de la clave. La tercera forma normal (3FN) exige que ningún atributo no clave dependa '
         'transitivamente de la clave primaria.'),
        ('Guía de SQL básico',
         'SELECT consulta filas y WHERE las filtra. Un LEFT JOIN conserva todas las filas de la tabla '
         'izquierda aunque no tengan coincidencia en la derecha. GROUP BY agrupa para usar COUNT o SUM.'),
        ('Calendario de Base de Datos I',
         'El primer parcial de Base de Datos I será el 15 de octubre en el aula B-204. '
         'El proyecto final se entrega el 28 de noviembre.'),
    ]),
    ('Programación', 'Ing. Mónica Alvarado', 'python, algoritmos', [
        ('Introducción a Python',
         'En Python las listas son mutables y las tuplas inmutables. La función len() devuelve la '
         'longitud de una secuencia y range(5) genera los números del 0 al 4.'),
        ('Ejercicios de bucles en Python',
         'Ejercicio 3: imprimir los números primos menores que 50 con un bucle for. '
         'Ejercicio 4: sumar los dígitos de un número con un bucle while.'),
        ('Rúbrica del proyecto de Programación I',
         'El proyecto de Programación I vale 30% de la nota: código 15%, documentación 10% y '
         'presentación 5%. Se trabaja en parejas.'),
    ]),
    ('Redes y Seguridad', 'Ing. Óscar Pineda', 'redes, osi, tcp/ip', [
        ('Resumen del modelo OSI',
         'El modelo OSI tiene 7 capas: física, enlace de datos, red, transporte, sesión, '
         'presentación y aplicación. TCP trabaja en la capa de transporte.'),
        ('Tabla de puertos comunes',
         'HTTP usa el puerto 80, HTTPS el 443, SSH el 22, DNS el 53 y SMTP el 25.'),
        ('Laboratorio de subredes',
         'Una red /26 tiene 64 direcciones y 62 hosts utilizables. La práctica de subredes se '
         'realiza en el laboratorio LAB-3 los viernes.'),
    ]),
    ('Inteligencia Artificial', 'Dr. Fernando Quintanilla', 'ia, machine learning', [
        ('Qué es el aprendizaje supervisado',
         'El aprendizaje supervisado entrena con ejemplos etiquetados. Sus dos tareas principales son '
         'la clasificación y la regresión.'),
        ('Glosario de inteligencia artificial',
         'Sobreajuste: el modelo memoriza los datos de entrenamiento y generaliza mal. '
         'Época: una pasada completa por el conjunto de entrenamiento.'),
        ('Charla de IA generativa en la UFG',
         'La charla "IA generativa en la educación" será el 9 de octubre a las 5:00 p.m. en el '
         'auditorio Gavidia. La entrada es libre.'),
    ]),
    ('Desarrollo Web', 'Ing. Gabriela Reyes', 'html, css, http', [
        ('Fundamentos de HTML y CSS',
         'Flexbox organiza elementos en una dimensión y CSS Grid en dos dimensiones. '
         'La etiqueta <main> marca el contenido principal de la página.'),
        ('Códigos de estado HTTP',
         '200 significa OK, 201 Created, 404 Not Found y 500 Internal Server Error. '
         'Los códigos 3xx indican redirecciones.'),
        ('Requisitos del proyecto web final',
         'El sitio del proyecto web final debe ser responsive, accesible y publicarse antes del '
         '20 de noviembre.'),
    ]),
    ('Sistemas Operativos', 'Ing. Héctor Campos', 'procesos, linux', [
        ('Planificación de procesos',
         'Round Robin asigna a cada proceso un quantum de tiempo fijo. SJF ejecuta primero el '
         'trabajo más corto.'),
        ('Comandos básicos de Linux',
         'ls lista archivos, cd cambia de carpeta, chmod cambia permisos y grep busca texto '
         'dentro de archivos.'),
        ('Interbloqueo y condiciones de Coffman',
         'Las cuatro condiciones de Coffman para un interbloqueo son: exclusión mutua, retención y '
         'espera, no expropiación y espera circular.'),
    ]),
    ('Ingeniería de Software', 'Ing. Claudia Mejía', 'scrum, requisitos', [
        ('Metodología Scrum',
         'Un sprint dura entre 1 y 4 semanas. Los roles de Scrum son Product Owner, Scrum Master '
         'y equipo de desarrollo.'),
        ('Requisitos funcionales y no funcionales',
         'Un requisito funcional describe lo que hace el sistema; uno no funcional describe cómo lo '
         'hace, por ejemplo rendimiento o seguridad.'),
        ('Plantilla de historia de usuario',
         'Formato: Como <rol> quiero <objetivo> para <beneficio>. Cada historia necesita criterios '
         'de aceptación.'),
    ]),
    ('Cálculo', 'Lic. Patricia Orellana', 'límites, derivadas', [
        ('Límites y continuidad',
         'Una función es continua en a si el límite en a existe y es igual a f(a).'),
        ('Tabla de derivadas básicas',
         'La derivada de sen(x) es cos(x), la de e^x es e^x y la de ln(x) es 1/x.'),
        ('Horario de tutorías de Cálculo II',
         'Las tutorías de Cálculo II son martes y jueves de 3:00 a 4:30 p.m. en el aula C-101 con '
         'la Lic. Patricia Orellana.'),
    ]),
    ('Estadística', 'Lic. Roberto Aguilar', 'media, distribución normal', [
        ('Medidas de tendencia central',
         'La media es el promedio, la mediana el valor central de los datos ordenados y la moda el '
         'valor más frecuente.'),
        ('La distribución normal',
         'En una distribución normal cerca del 68% de los datos está a una desviación estándar de '
         'la media y el 95% a dos.'),
        ('Datos de práctica de estadística',
         'Notas de práctica: 7.5, 8.0, 6.5, 9.0 y 7.0. La media de estas notas es 7.6.'),
    ]),
    ('Física', 'Ing. Mauricio Lemus', 'newton, cinemática', [
        ('Leyes de Newton',
         'Primera ley: inercia. Segunda ley: F = m·a. Tercera ley: a toda acción corresponde una '
         'reacción igual y opuesta.'),
        ('Fórmulas de cinemática',
         'v = v0 + a·t y x = x0 + v0·t + a·t²/2. La gravedad se aproxima a 9.8 m/s².'),
        ('Práctica del péndulo simple',
         'El periodo de un péndulo simple es T = 2*pi*raiz(L/g). La práctica se entrega en el '
         'laboratorio de física F-2.'),
    ]),
    ('Química', 'Lic. Karla Benítez', 'tabla periódica, enlaces', [
        ('Tabla periódica básica',
         'El número atómico del carbono es 6, el del nitrógeno 7 y el del oxígeno 8.'),
        ('Enlaces químicos',
         'El enlace iónico transfiere electrones y el covalente los comparte. El agua tiene enlaces '
         'covalentes polares.'),
        ('Normas del laboratorio de química',
         'En el laboratorio de química es obligatorio usar bata y gafas de seguridad. Está prohibido '
         'comer o beber.'),
    ]),
    ('Álgebra', 'Lic. Sergio Durán', 'matrices, sistemas', [
        ('Operaciones con matrices',
         'Dos matrices se pueden multiplicar si el número de columnas de la primera es igual al '
         'número de filas de la segunda.'),
        ('Sistemas de ecuaciones lineales',
         'El método de Gauss transforma el sistema en una matriz escalonada para resolverlo por '
         'sustitución hacia atrás.'),
        ('Ejercicios de determinantes',
         'El determinante de la matriz [[2, 1], [3, 4]] es 5. Una matriz con determinante 0 no '
         'tiene inversa.'),
    ]),
    ('Diseño Gráfico', 'Lic. Daniela Portillo', 'color, tipografía', [
        ('Teoría del color',
         'Los colores primarios de la luz (RGB) son rojo, verde y azul; los del pigmento (CMYK) son '
         'cian, magenta y amarillo.'),
        ('Introducción a la tipografía',
         'Las fuentes serif tienen remates en los trazos; las sans serif no. Para pantallas se '
         'suelen preferir las sans serif.'),
        ('Concurso de afiches UFG',
         'Las inscripciones para el concurso de afiches cierran el 30 de octubre. El primer lugar '
         'recibe $150.'),
    ]),
    ('Diseño UX/UI', 'Lic. Tomás Cornejo', 'usabilidad, nielsen', [
        ('Heurísticas de Nielsen',
         'La primera heurística de Nielsen es la visibilidad del estado del sistema. Son 10 '
         'heurísticas en total.'),
        ('Wireframes de baja fidelidad',
         'Un wireframe de baja fidelidad usa cajas y líneas para definir la estructura sin colores '
         'ni imágenes.'),
        ('Ley de Fitts',
         'La ley de Fitts dice que el tiempo para alcanzar un objetivo depende de su distancia y de '
         'su tamaño.'),
    ]),
    ('Arquitectura', 'Arq. Elena Villalta', 'bioclimática, planos', [
        ('Arquitectura bioclimática',
         'En El Salvador, orientar las fachadas largas al norte y al sur reduce la ganancia solar. '
         'La ventilación cruzada ayuda a enfriar.'),
        ('Materiales sostenibles',
         'El bambú y el adobe son materiales de baja huella de carbono usados en la región.'),
        ('Escalas en planos arquitectónicos',
         'Una escala 1:100 significa que 1 cm en el plano equivale a 100 cm en la realidad.'),
    ]),
    ('Contabilidad', 'Lic. Mario Escobar', 'balance, partida doble', [
        ('Partida doble',
         'En la partida doble todo cargo tiene un abono por el mismo valor. Activo = Pasivo + '
         'Capital.'),
        ('Estados financieros básicos',
         'Los estados financieros básicos son el balance general, el estado de resultados y el '
         'flujo de efectivo.'),
        ('Ejercicio de depreciación',
         'Un equipo de $3,000 con vida útil de 5 años se deprecia $600 por año en línea recta.'),
    ]),
    ('Marketing', 'Lic. Silvia Rosales', '4p, segmentación', [
        ('Las 4P del marketing',
         'Las 4P del marketing son producto, precio, plaza y promoción.'),
        ('Segmentación de mercado',
         'La segmentación divide el mercado en grupos por criterios demográficos, geográficos, '
         'psicográficos o conductuales.'),
        ('Caso de estudio de marca local',
         'La marca ficticia "Pupusa Express" aumentó sus ventas un 18% tras una campaña en redes '
         'sociales.'),
    ]),
    ('Derecho Civil', 'Lic. Raúl Castaneda', 'contratos, personas', [
        ('Elementos del contrato',
         'Los elementos esenciales del contrato son consentimiento, objeto y causa.'),
        ('Capacidad jurídica',
         'La capacidad de goce la tienen todas las personas; la capacidad de ejercicio se adquiere '
         'con la mayoría de edad, a los 18 años.'),
        ('Calendario de Derecho Civil I',
         'El examen de Derecho Civil I será el 22 de octubre en el aula D-305.'),
    ]),
    ('Psicología', 'Lic. Irene Guardado', 'piaget, memoria', [
        ('Etapas de Piaget',
         'Las etapas del desarrollo de Piaget son: sensoriomotora, preoperacional, operaciones '
         'concretas y operaciones formales.'),
        ('Tipos de memoria',
         'La memoria de trabajo mantiene unos 7 elementos a la vez; la memoria a largo plazo tiene '
         'capacidad casi ilimitada.'),
        ('Prácticas de psicología clínica',
         'Las prácticas de psicología clínica empiezan en el sexto ciclo y requieren 120 horas.'),
    ]),
    ('Anatomía', 'Dra. Beatriz Navas', 'huesos, corazón', [
        ('Sistema óseo',
         'El cuerpo humano adulto tiene 206 huesos. El fémur es el hueso más largo.'),
        ('El corazón',
         'El corazón tiene cuatro cavidades: dos aurículas y dos ventrículos.'),
        ('Laboratorio de anatomía',
         'El laboratorio de anatomía abre de lunes a viernes de 7:00 a.m. a 12:00 m. en el edificio '
         'de Ciencias de la Salud.'),
    ]),
    ('Inglés', 'Lic. Karen Ayala', 'gramática, tiempos', [
        ('Present perfect',
         'El present perfect se forma con have/has + participio: "I have finished my homework".'),
        ('Verbos irregulares frecuentes',
         'go-went-gone, see-saw-seen, write-wrote-written, take-took-taken.'),
        ('Examen de suficiencia de inglés',
         'El examen de suficiencia de inglés es el 5 de noviembre; se aprueba con 80 puntos de 100.'),
    ]),
    ('Ingeniería Civil', 'Ing. Arturo Peña', 'concreto, estructuras', [
        ('Resistencia del concreto',
         'Un concreto de 210 kg/cm² es el más usado en viviendas. Alcanza su resistencia de diseño a '
         'los 28 días.'),
        ('Tipos de cargas en estructuras',
         'Las cargas muertas son permanentes (peso propio) y las cargas vivas son variables (personas, '
         'muebles).'),
        ('Visita técnica a obra',
         'La visita técnica a la obra del puente Las Delicias será el 17 de octubre; salida a las '
         '7:00 a.m. desde el portón principal.'),
    ]),
]

# Grupos: (clave, nombre, descripción, dueño, miembros, visible_para_demo, documentos, tareas, mensajes)
GRUPOS = [
    ('bd', 'Proyecto Base de Datos II', 'Equipo del proyecto final de Base de Datos II',
     'carlos', ['ana', 'javier'], True,
     [('Modelo entidad-relación del proyecto',
       'La base de datos del proyecto tiene 8 tablas: estudiantes, cursos, inscripciones, docentes, '
       'aulas, horarios, notas y pagos.'),
      ('Reparto de tareas del equipo',
       'Ana diseña las vistas, Carlos los procedimientos almacenados y Javier las pruebas.'),
      ('Acta de la reunión 1',
       'Se acordó usar PostgreSQL 16 y entregar el primer avance el 18 de octubre.'),
      ('Datos de prueba del proyecto',
       'Tabla cursos: BD2-01 Base de Datos II 4 UV; PRG-02 Programación II 4 UV; RED-01 Redes I 3 UV.')],
     [('Diseñar el modelo ER', 'Primera versión del diagrama'), ('Crear procedimientos almacenados', '')],
     [('carlos', '¡Bienvenidos al grupo! Ya subí el modelo ER.'),
      ('ana', 'Gracias Carlos, empiezo con las vistas hoy.'),
      ('javier', 'Yo preparo los datos de prueba para el viernes.'),
      ('carlos', 'Recuerden: primer avance el 18 de octubre.')]),
    ('arq', 'Tesis Arquitectura Sostenible', 'Seguimiento de la tesis de arquitectura bioclimática',
     'lucia', ['javier'], True,
     [('Propuesta de tesis',
       'Centro comunitario en Santa Tecla con ventilación cruzada, techos verdes y captación de agua lluvia.'),
      ('Cronograma de la tesis',
       'Entrega del anteproyecto el 30 de octubre y defensa prevista para el 12 de diciembre.'),
      ('Bibliografía de la tesis',
       'Olgyay, "Arquitectura y clima"; Neila, "Arquitectura bioclimática en un entorno sostenible".'),
      ('Presupuesto preliminar',
       'Costo estimado del centro comunitario: $185,000. Estructura 45%, acabados 30%, instalaciones 25%.')],
     [('Terminar planta arquitectónica', 'Escala 1:100'), ('Revisar bibliografía', '')],
     [('lucia', 'Subí la propuesta, ¿la revisan?'),
      ('javier', 'Me gusta lo de los techos verdes. Te ayudo con el cálculo estructural.'),
      ('lucia', '¡Genial! La defensa sería el 12 de diciembre.')]),
    ('calc', 'Estudio Cálculo II', 'Grupo de estudio para los parciales de Cálculo II',
     'ana', ['diego', 'javier', 'valeria'], True,
     [('Resumen de integración por partes',
       'Integral de u dv = u·v - integral de v du. Para elegir u se usa la regla LIATE: logarítmicas, '
       'inversas, algebraicas, trigonométricas y exponenciales.'),
      ('Ejercicios resueltos de series',
       'La serie geométrica de razón 1/2 que empieza en 1 converge a 2. La serie armónica diverge.'),
      ('Plan de estudio para el parcial',
       'Nos reunimos el sábado 11 de octubre a las 9:00 a.m. en la biblioteca central, sala 3.'),
      ('Formulario de integrales',
       'Integral de 1/x = ln|x| + C. Integral de e^x = e^x + C. Integral de cos(x) = sen(x) + C.')],
     [('Resolver guía 3', 'Ejercicios 1 al 10'), ('Repasar series', '')],
     [('ana', '¿Nos vemos el sábado en la biblioteca?'),
      ('diego', 'Sí, a las 9 me queda bien.'),
      ('valeria', 'Llevo los ejercicios de series resueltos.'),
      ('javier', 'Yo no entiendo LIATE, ¿me explican el sábado?'),
      ('ana', 'Claro, subí un resumen al grupo.')]),
    ('emp', 'Emprendimiento UFG', 'Plan de negocio para la feria de emprendimiento',
     'diego', ['sofia', 'andres'], False,
     [('Plan de negocio CaféTec',
       'CaféTec es una cafetería dentro del campus con inversión inicial de $12,500 y punto de '
       'equilibrio en el mes 14.'),
      ('Lienzo Canvas de CaféTec',
       'Segmento: estudiantes y docentes. Propuesta de valor: café de especialidad salvadoreño a $1.50.'),
      ('Estudio de mercado CaféTec',
       'El 62% de los 150 encuestados compraría café dentro del campus al menos tres veces por semana.'),
      ('Proyección financiera CaféTec',
       'Ventas año 1: $48,000. Costos año 1: $39,500. Utilidad estimada: $8,500.')],
     [('Preparar pitch de 3 minutos', 'Para la feria'), ('Cotizar máquina de espresso', '')],
     [('diego', 'Equipo, la feria es el 25 de octubre.'),
      ('sofia', 'Puedo hacer la encuesta de satisfacción.'),
      ('andres', 'Reviso el tema de permisos municipales.')]),
    ('psi', 'Club de Psicología Cognitiva', 'Lecturas y debates sobre psicología cognitiva',
     'sofia', ['valeria', 'andres'], False,
     [('Sesgos cognitivos frecuentes',
       'El sesgo de confirmación lleva a buscar sólo la información que confirma lo que ya creemos.'),
      ('Lecturas del club',
       'Este ciclo leemos "Pensar rápido, pensar despacio" de Daniel Kahneman, un capítulo por semana.'),
      ('Encuesta de hábitos de sueño',
       'De 80 estudiantes encuestados, el 71% duerme menos de 6 horas en semana de parciales.'),
      ('Agenda del club',
       'Las sesiones del club son los miércoles a las 6:00 p.m. en el aula P-12.')],
     [('Leer capítulo 3', 'Kahneman'), ('Diseñar nueva encuesta', '')],
     [('sofia', 'El miércoles comentamos el capítulo 3.'),
      ('valeria', '¡Me encantó lo del sistema 1 y sistema 2!'),
      ('andres', 'Traigo ejemplos de sesgos en juicios legales.')]),
]

# Archivos privados (14): (dueño, título, texto)
PRIVADOS = [
    ('ana', 'Mi horario del ciclo', 'Lunes y miércoles: Base de Datos II 7:00 a.m. Martes: Cálculo II. Mi casillero es el 4471.'),
    ('ana', 'Pendientes personales', 'Pagar la cuota de octubre antes del día 10. Renovar el carné de biblioteca.'),
    ('carlos', 'Ideas para la tesis', 'Tema posible: sistema de recomendación de cursos con filtrado colaborativo.'),
    ('carlos', 'Contraseñas de laboratorio (ficticias)', 'Usuario del servidor de pruebas del LAB-3: alumno07. Clave de práctica: lab3-practica.'),
    ('lucia', 'Bocetos de fachada', 'Fachada norte con parasoles verticales de madera y celosías de barro.'),
    ('lucia', 'Presupuesto de maqueta', 'Cartón pluma $12, madera balsa $18, pintura $7. Total $37.'),
    ('diego', 'Notas de la reunión con inversionista', 'El inversionista ficticio ofrece $5,000 a cambio del 10% de CaféTec.'),
    ('diego', 'Lista de proveedores', 'Café: Finca El Socorro. Vasos: Distribuidora Central.'),
    ('sofia', 'Diario de lecturas', 'Terminé el capítulo 2 de Kahneman. Me sorprendió el efecto ancla.'),
    ('javier', 'Cálculos de viga', 'Viga de 6 m con carga distribuida de 2 t/m: momento máximo 9 t·m.'),
    ('javier', 'Checklist de visita a obra', 'Llevar casco, chaleco, botas y libreta de campo.'),
    ('valeria', 'Resumen de fisiología', 'La frecuencia cardiaca normal en reposo es de 60 a 100 latidos por minuto.'),
    ('andres', 'Borrador de ensayo legal', 'Tema: la mediación como alternativa a los juicios civiles.'),
    ('andres', 'Fechas de audiencias simuladas', 'Audiencia simulada del consultorio jurídico: 3 de noviembre, sala 2.'),
]

# Anotaciones de texto: (dueño, título, contenido)
NOTAS = [
    ('ana', 'Recordatorio parcial', 'Estudiar normalización y JOINs para el parcial del 15 de octubre.'),
    ('ana', 'Idea de app', 'App para reservar cubículos de la biblioteca.'),
    ('carlos', 'Para el equipo', 'Revisar índices de la tabla inscripciones.'),
    ('lucia', 'Referentes', 'Ver obras de Anna Heringer con tierra y bambú.'),
    ('diego', 'Pitch', 'Abrir con el dato del 62% del estudio de mercado.'),
    ('sofia', 'Para el club', 'Proponer debate sobre el efecto Dunning-Kruger.'),
    ('javier', 'Materiales', 'Comparar precios de varilla de 3/8 en dos ferreterías.'),
    ('valeria', 'Anatomía', 'Repasar las cavidades del corazón y el sistema óseo.'),
    ('andres', 'Lecturas', 'Leer el capítulo de obligaciones del Código Civil.'),
    ('valeria', 'Mi PIN del casillero', 'El PIN ficticio de mi casillero es 8823 (dato privado para probar permisos).'),
]

CHAT_GLOBAL = [
    ('ana', '¡Hola a todos! ¿Alguien tiene apuntes de Base de Datos I?'),
    ('carlos', 'Sí, en la Biblioteca hay unos de normalización.'),
    ('lucia', '¿Alguien va a la charla de IA generativa del 9 de octubre?'),
    ('diego', 'Yo voy, me interesa para el emprendimiento.'),
    ('sofia', 'Recuerden que el club de psicología se reúne los miércoles.'),
    ('javier', '¿Quién más va a la visita técnica del puente Las Delicias?'),
    ('valeria', 'El laboratorio de anatomía cierra a mediodía, ojo.'),
    ('andres', 'Busco compañeros para practicar inglés antes del examen de suficiencia.'),
    ('ana', 'Andrés, yo me apunto.'),
    ('carlos', 'Subí una guía de SQL básico por si a alguien le sirve.'),
    ('lucia', 'El concurso de afiches cierra el 30 de octubre, ¡anímense!'),
    ('diego', 'Gracias por la info, Lucía.'),
]

# Contactos aceptados entre usuarios de prueba: (quien_envía, quien_acepta)
CONTACTOS = [
    ('ana', 'carlos'), ('ana', 'javier'), ('ana', 'diego'), ('carlos', 'javier'),
    ('lucia', 'javier'), ('diego', 'sofia'), ('diego', 'andres'), ('sofia', 'valeria'),
    ('sofia', 'andres'), ('valeria', 'andres'),
]
# Solicitudes pendientes para el usuario demo (se aceptan desde la interfaz).
PENDIENTES_DEMO = ['ana', 'carlos', 'lucia']

MENSAJES_DIRECTOS = [
    ('ana', 'carlos', '¿Me pasas el enlace del repositorio del proyecto?'),
    ('carlos', 'ana', 'Claro, te lo mando en un rato.'),
    ('ana', 'carlos', '¡Gracias!'),
    ('diego', 'sofia', '¿Puedes ayudarme con la encuesta de CaféTec?'),
    ('sofia', 'diego', 'Sí, el jueves la tengo lista.'),
    ('lucia', 'javier', '¿Revisaste el cálculo estructural?'),
    ('javier', 'lucia', 'Casi, mañana te lo envío.'),
    ('valeria', 'andres', '¿Vienes al club el miércoles?'),
    ('andres', 'valeria', 'Sí, llevo los ejemplos de sesgos.'),
]


# --- Utilidades ----------------------------------------------------------------

def slug(texto, sep='-'):
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', sep, texto.lower()).strip(sep)


FORMATOS = ['pdf', 'docx', 'pptx', 'pdf', 'xlsx', 'docx', 'pptx', 'txt', 'pdf', 'md']


def formato(i):
    """Reparte los 100 documentos: 30 pdf, 20 docx, 20 pptx, 10 xlsx, 10 txt y 10 md/csv."""
    ext = FORMATOS[i % len(FORMATOS)]
    if ext == 'md' and (i // len(FORMATOS)) % 2:
        ext = 'csv'
    return ext


def _pdf_texto(s):
    # Helvetica con WinAnsiEncoding: sólo caracteres de cp1252.
    s = s.encode('cp1252', 'replace')
    return s.replace(b'\\', b'\\\\').replace(b'(', b'\\(').replace(b')', b'\\)')


def crear_pdf(ruta, titulo, texto):
    """PDF mínimo de una página escrito a mano (sin dependencias)."""
    lineas = textwrap.wrap(texto, 85)
    flujo = [b'BT /F1 16 Tf 72 760 Td (' + _pdf_texto(titulo) + b') Tj /F1 11 Tf 0 -30 Td']
    for linea in lineas:
        flujo.append(b'(' + _pdf_texto(linea) + b') Tj 0 -16 Td')
    flujo.append(b'ET')
    contenido = b'\n'.join(flujo)
    objetos = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
        b'/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
        b'<< /Length ' + str(len(contenido)).encode() + b' >>\nstream\n' + contenido + b'\nendstream',
    ]
    salida = io.BytesIO()
    salida.write(b'%PDF-1.4\n')
    posiciones = []
    for n, obj in enumerate(objetos, 1):
        posiciones.append(salida.tell())
        salida.write(f'{n} 0 obj\n'.encode() + obj + b'\nendobj\n')
    xref = salida.tell()
    salida.write(f'xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n'.encode())
    for p in posiciones:
        salida.write(f'{p:010d} 00000 n \n'.encode())
    salida.write(f'trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    ruta.write_bytes(salida.getvalue())


def crear_docx(ruta, titulo, texto):
    import docx
    d = docx.Document()
    d.add_heading(titulo, level=1)
    d.add_paragraph(texto)
    d.add_paragraph('Material de prueba de UFG Knowledge Hub.')
    d.save(ruta)


def crear_pptx(ruta, titulo, texto):
    from pptx import Presentation
    p = Presentation()
    portada = p.slides.add_slide(p.slide_layouts[0])
    portada.shapes.title.text = titulo
    portada.placeholders[1].text = 'UFG Knowledge Hub - material de prueba'
    diapositiva = p.slides.add_slide(p.slide_layouts[1])
    diapositiva.shapes.title.text = 'Contenido'
    cuerpo = diapositiva.placeholders[1].text_frame
    frases = [f.strip() for f in re.split(r'(?<=\.)\s+', texto) if f.strip()]
    cuerpo.text = frases[0]
    for frase in frases[1:]:
        cuerpo.add_paragraph().text = frase
    p.save(ruta)


def crear_xlsx(ruta, titulo, texto):
    from openpyxl import Workbook
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Datos'
    hoja.append([titulo])
    for frase in re.split(r'(?<=\.)\s+', texto):
        if frase.strip():
            hoja.append([frase.strip()])
    wb.save(ruta)


def crear_archivo(carpeta, titulo, texto, ext):
    ruta = carpeta / f'{slug(titulo)}.{ext}'
    if ext == 'pdf':
        crear_pdf(ruta, titulo, texto)
    elif ext == 'docx':
        crear_docx(ruta, titulo, texto)
    elif ext == 'pptx':
        crear_pptx(ruta, titulo, texto)
    elif ext == 'xlsx':
        crear_xlsx(ruta, titulo, texto)
    elif ext == 'md':
        ruta.write_text(f'# {titulo}\n\n{texto}\n', encoding='utf-8')
    elif ext == 'csv':
        with open(ruta, 'w', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh)
            w.writerow(['titulo', 'contenido'])
            w.writerow([titulo, texto])
    else:
        ruta.write_text(f'{titulo}\n\n{texto}\n', encoding='utf-8')
    return ruta


# --- Cliente de la API ---------------------------------------------------------

class Api:
    def __init__(self, url):
        self.url = url.rstrip('/')
        self.tokens = {}
        self.ids = {}

    def _check(self, r, esperado=(200, 201)):
        if r.status_code not in esperado:
            raise RuntimeError(f'{r.request.method} {r.url} -> {r.status_code}: {r.text[:300]}')
        return r.json() if r.content else {}

    def h(self, clave):
        return {'Authorization': f'Bearer {self.tokens[clave]}'}

    def post(self, clave, ruta, **kw):
        return self._check(requests.post(self.url + ruta, headers=self.h(clave), timeout=60, **kw))

    def get(self, clave, ruta, **kw):
        return self._check(requests.get(self.url + ruta, headers=self.h(clave), timeout=60, **kw))

    def put(self, clave, ruta, **kw):
        return self._check(requests.put(self.url + ruta, headers=self.h(clave), timeout=60, **kw))

    def registrar_y_entrar(self, clave, nombre, color):
        if not crear_cuenta_de_prueba(email_de(clave), nombre, color):
            return False
        datos = self._check(requests.post(self.url + '/api/login', json={
            'email': email_de(clave), 'password': TEST_PASSWORD}, timeout=60))
        self.tokens[clave] = datos['token']
        self.ids[clave] = datos['user']['id']
        return True

    def subir(self, clave, ruta_api, archivo, campos):
        with open(archivo, 'rb') as fh:
            return self.post(clave, ruta_api, data=campos, files={'file': (archivo.name, fh)})


_app_local = None


def crear_cuenta_de_prueba(email, nombre, color):
    """Crea la cuenta en la base de datos local. Devuelve False si ya existía."""
    global _app_local
    from hub import create_app
    from hub.extensions import bcrypt, db
    from hub.models import User

    if _app_local is None:
        _app_local = create_app({'BACKGROUND_WORKERS': False})
    with _app_local.app_context():
        if User.query.filter(db.func.lower(User.email) == email.lower()).first():
            return False
        db.session.add(User(name=nombre, email=email, profile_color=color,
                            password=bcrypt.generate_password_hash(TEST_PASSWORD).decode('utf-8')))
        db.session.commit()
    return True


def main():
    parser = argparse.ArgumentParser(description='Crea datos de prueba a través de la API.')
    parser.add_argument('--url', default='http://127.0.0.1:5001', help='URL de la app en marcha')
    args = parser.parse_args()

    api = Api(args.url)
    try:
        requests.get(api.url + '/api/health', timeout=5)
    except requests.RequestException:
        sys.exit(f'No se pudo conectar con {api.url}. Arranca la app con "python app.py".')

    carpeta = Path(tempfile.mkdtemp(prefix='ufg_seed_'))
    try:
        # Usuarios
        for clave, nombre, carrera, semestre, color in USUARIOS:
            if not api.registrar_y_entrar(clave, nombre, color):
                sys.exit(f'El usuario {email_de(clave)} ya existe: los datos de prueba ya se crearon. '
                         'Para empezar de cero: python manage.py db --reset')
            api.post(clave, '/api/profile', json={'name': nombre, 'career': carrera, 'semester': semestre})
        print(f'Usuarios: {len(USUARIOS)}')

        demo = api.get('ana', '/api/contacts/search', params={'q': DEMO_EMAIL})['users']
        demo_id = demo[0]['id'] if demo else None

        n_doc = 0
        claves = [u[0] for u in USUARIOS]

        # Biblioteca (66)
        for categoria, autor, etiquetas, docs in BIBLIOTECA:
            for titulo, texto in docs:
                ext = formato(n_doc)
                archivo = crear_archivo(carpeta, titulo, texto, ext)
                api.subir(claves[n_doc % len(claves)], '/api/resources', archivo, {
                    'title': titulo, 'author': autor, 'category': categoria, 'tags': etiquetas})
                n_doc += 1
        print(f'Biblioteca: {n_doc} documentos')

        # Grupos (20 documentos)
        for _, nombre, descripcion, dueno, miembros, visible_demo, docs, tareas, mensajes in GRUPOS:
            gid = api.post(dueno, '/api/groups', json={'name': nombre, 'description': descripcion})['group_id']
            for m in miembros:
                api.post(dueno, f'/api/groups/{gid}/invite', json={'user_id': api.ids[m]})
            if visible_demo and demo_id:
                api.post(dueno, f'/api/groups/{gid}/invite', json={'user_id': demo_id})
            integrantes = [dueno] + miembros
            for i, (titulo, texto) in enumerate(docs):
                ext = formato(n_doc)
                archivo = crear_archivo(carpeta, titulo, texto, ext)
                api.subir(integrantes[i % len(integrantes)], f'/api/groups/{gid}/resources', archivo, {
                    'title': titulo, 'description': f'Documento del grupo {nombre}'})
                n_doc += 1
            for titulo, desc in tareas:
                api.post(dueno, f'/api/groups/{gid}/tasks', json={
                    'title': titulo, 'description': desc, 'due_date': '2026-11-15T23:59:00'})
            for autor_msg, texto in mensajes:
                api.post(autor_msg, f'/api/groups/{gid}/messages', json={'message': texto})
        print(f'Grupos: {len(GRUPOS)} (documentos acumulados: {n_doc})')

        # Archivos privados (14)
        for dueno, titulo, texto in PRIVADOS:
            ext = formato(n_doc)
            archivo = crear_archivo(carpeta, titulo, texto, ext)
            api.subir(dueno, '/api/private-resources', archivo, {'title': titulo})
            n_doc += 1
        print(f'Documentos en total: {n_doc}')

        for dueno, titulo, contenido in NOTAS:
            api.post(dueno, '/api/private-resources', json={'title': titulo, 'content': contenido})
        print(f'Anotaciones de texto: {len(NOTAS)}')

        # Contactos
        for de, para in CONTACTOS:
            api.post(de, '/api/contacts/request', json={'contact_id': api.ids[para]})
            solicitud = next(s for s in api.get(para, '/api/contacts/requests')['requests']
                             if s['user_id'] == api.ids[de])
            api.put(para, f'/api/contacts/requests/{solicitud["id"]}/accept')
        if demo_id:
            for de in PENDIENTES_DEMO:
                api.post(de, '/api/contacts/request', json={'contact_id': demo_id})
        print(f'Contactos: {len(CONTACTOS)} aceptados, {len(PENDIENTES_DEMO) if demo_id else 0} pendientes para demo')

        for de, para, texto in MENSAJES_DIRECTOS:
            api.post(de, '/api/direct-messages', json={'receiver_id': api.ids[para], 'message': texto})
        print(f'Mensajes directos: {len(MENSAJES_DIRECTOS)}')

        for autor_msg, texto in CHAT_GLOBAL:
            api.post(autor_msg, '/api/chat/messages', json={'message': texto})
        print(f'Chat global: {len(CHAT_GLOBAL)} mensajes')

        print('\nCuentas de prueba (contraseña: TEST_PASSWORD en este script):')
        for clave, nombre, *_ in USUARIOS:
            print(f'  {nombre:<18} {email_de(clave)}')
        if not demo_id:
            print(f'\nAviso: no existe {DEMO_EMAIL}; no se le añadieron grupos ni solicitudes.')
    finally:
        shutil.rmtree(carpeta, ignore_errors=True)


if __name__ == '__main__':
    main()
